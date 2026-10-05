# Công thức nháp v1 để đoán nhãn Datangon (Vào ngay / Đang hot / Né ra / Đang giảm / Theo dõi).
# Chạy:  python3 cong_thuc_v1.py     (cần: pandas numpy openpyxl)
import json, os, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
V_NGUONG = 7.25      # "video ít" nếu mức video trung bình < số này (thang 0-20)
D_NGUONG = 3000      # đơn 30 ngày từ số này trở lên coi là "đơn lớn"

def doc():
    j = json.load(open(os.path.join(HERE, "du_lieu", "datangon-100-ngach.json"), encoding="utf-8"))
    rows = []
    for r in j:
        tv = [24 - float(x) for x in r["trend_video_ys"].split(";")] if r["trend_video_ys"] else []
        to = [24 - float(x) for x in r["trend_don_ys"].split(";")] if r["trend_don_ys"] else []
        rows.append(dict(
            nganh=r["nganh_hang"], ngach=r["nganh_hang_chi_tiet"], that=r["nhan"], id=r["id_sp"],
            don=r["don30n"] if r["don30n"] not in ("", None) else None,
            mui_ten=(r["mui_ten"] or "").replace("(màu)", ""),
            video_tb=float(np.mean(tv)) if len(tv) == 5 else None,
            don_cuoi=to[-1] if len(to) == 5 else None,
            da_ban=r["da_ban"], video_moi=r["video_moi_pct"], link=r["link_sp"]))
    return pd.DataFrame(rows)

def doan(r):
    """Luật: nhìn 3 thứ  -  mũi tên đơn (tăng/giảm), mức video trung bình, số đơn 30 ngày."""
    if r.don is None or pd.isna(r.don):
        return "Đang giảm"                      # không có đơn nào trong 30 ngày
    tang = r.mui_ten == "tăng"
    it_video = (r.video_tb is not None) and (not pd.isna(r.video_tb)) and r.video_tb < V_NGUONG
    if tang:
        return "Vào ngay" if it_video else "Đang hot"
    if not it_video:                            # đơn giảm nhưng nhiều video
        return "Đang hot" if r.don >= D_NGUONG else "Né ra"
    return "Đang giảm"                          # đơn giảm và ít video

if __name__ == "__main__":
    df = doc()
    df["doan"] = df.apply(doan, axis=1)
    df["dung"] = df.that == df.doan
    chinh = df.that.isin(["Vào ngay", "Đang hot", "Né ra"])
    print("Tổng số dòng:", len(df))
    print("Đoán đúng (cả 5 nhãn): %.1f%%" % (100 * df.dung.mean()))
    print("Đoán đúng (3 nhãn chính): %.1f%%" % (100 * df.dung[chinh].mean()))
    print(pd.crosstab(df.that, df.doan))

    wb = Workbook(); ws = wb.active; ws.title = "Thật và đoán"
    cols = ["Ngành lớn", "Ngách con", "Nhãn THẬT", "Nhãn ĐOÁN", "Đúng?", "Đơn 30N", "Đơn đang", "Mức video TB (0-20)", "Đã bán", "Link"]
    ws.append(cols)
    for _, r in df.sort_values(["dung", "nganh", "ngach"]).iterrows():
        ws.append([r.nganh, r.ngach, r.that, r.doan, "Đúng" if r.dung else "SAI",
                   None if pd.isna(r.don) else int(r.don), r.mui_ten,
                   None if pd.isna(r.video_tb) else round(r.video_tb, 1), int(r.da_ban), r.link])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    for row in ws.iter_rows(min_row=2):
        if row[4].value == "SAI":
            for c in row: c.fill = PatternFill("solid", fgColor="FED7D7")
    for i, w in enumerate([26, 28, 12, 12, 8, 10, 9, 14, 11, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(os.path.join(HERE, "so-sanh-nhan-that-va-doan.xlsx"))
    print("Đã tạo so-sanh-nhan-that-va-doan.xlsx")
